#!/usr/bin/env python3
"""Read-only audit of Pihe, Shihe, and Hangbuhe manual pond totals."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SURVEY_BOOK = ROOT.parent / "2024年-2026年蓄水遥感人工数据.xlsx"
LEDGER = ROOT / "diagnostics/pond_stage_ledger.xlsx"
TARGETS = (("2024年人工", "2024-06-01"), ("2025年人工", "2025-05-21"))
REGIONS = {
    "pihe": "淠河",
    "shihe": "史河",
    "hangbuhe": "杭埠河",
}
OUTPUT_ROOT = ROOT / "diagnostics/results/cross_region_manual_totals"


def numeric(value) -> float:
    if value in (None, ""):
        return 0.0
    return float(value)


def find_date_column(ws, target_date: str) -> tuple[int, int]:
    # Merged cells may expose ``cell.column=None`` and some workbooks may not
    # report a usable max_column.  enumerate() provides a guaranteed integer.
    for column, cell in enumerate(ws[2], start=1):
        if str(cell.value).strip() == target_date:
            pond_col = column + 3
            if ws.cell(3, pond_col).value != "塘坝":
                raise RuntimeError(f"{ws.title} {target_date}: expected pond header")
            return column, pond_col
    raise RuntimeError(f"{ws.title}: date not found {target_date}")


def find_region_rows(ws, token: str) -> tuple[int, int, str]:
    matches = []
    max_row = int(ws.max_row or 0)
    for row in range(1, max_row + 1):
        value = ws.cell(row, 1).value
        if value not in (None, "") and token in str(value):
            matches.append((row, str(value)))
    if len(matches) != 1:
        raise RuntimeError(f"{ws.title}: expected one region block containing {token}, found {matches}")
    start, label = matches[0]
    end = max_row
    for row in range(start + 1, max_row + 1):
        if ws.cell(row, 1).value not in (None, ""):
            end = row - 1
            break
    totals = [row for row in range(start, end + 1) if ws.cell(row, 2).value == "合计"]
    if len(totals) != 1:
        raise RuntimeError(f"{ws.title} {label}: expected one total row, found {totals}")
    return start, totals[0], label


def ledger_values() -> dict[tuple[str, str], float]:
    workbook = openpyxl.load_workbook(LEDGER, read_only=True, data_only=True)
    ws = workbook["阶段台账"]
    max_column = int(ws.max_column or 0)
    max_row = int(ws.max_row or 0)
    headers = {ws.cell(1, col).value: col for col in range(1, max_column + 1)}
    values = {}
    for row in range(2, max_row + 1):
        key = (str(ws.cell(row, headers["date"]).value), str(ws.cell(row, headers["region"]).value))
        values[key] = numeric(ws.cell(row, headers["manual_pond_10k_m3"]).value)
    workbook.close()
    return values


def main() -> int:
    if not SURVEY_BOOK.exists():
        raise FileNotFoundError(SURVEY_BOOK)
    formulas = openpyxl.load_workbook(SURVEY_BOOK, read_only=False, data_only=False)
    values = openpyxl.load_workbook(SURVEY_BOOK, read_only=False, data_only=True)
    ledger = ledger_values()
    records = []
    for sheet, target_date in TARGETS:
        ws_value = values[sheet]
        ws_formula = formulas[sheet]
        date_col, pond_col = find_date_column(ws_value, target_date)
        for region, token in REGIONS.items():
            start, total_row, source_label = find_region_rows(ws_value, token)
            component_rows = list(range(start, total_row))
            component_sum = sum(
                numeric(ws_value.cell(row, pond_col).value) for row in component_rows
            )
            source_total = numeric(ws_value.cell(total_row, pond_col).value)
            ledger_value = ledger.get((target_date, region))
            records.append({
                "sheet": sheet,
                "date": target_date,
                "region": region,
                "source_region_label": source_label,
                "date_header_cell": ws_value.cell(2, date_col).coordinate,
                "pond_header_cell": ws_value.cell(3, pond_col).coordinate,
                "total_cell": ws_value.cell(total_row, pond_col).coordinate,
                "total_raw_formula_or_value": ws_formula.cell(total_row, pond_col).value,
                "source_total_wan_m3": source_total,
                "component_row_start": start,
                "component_row_end": total_row - 1,
                "component_sum_wan_m3": component_sum,
                "component_minus_total_wan_m3": component_sum - source_total,
                "component_sum_matches_total": abs(component_sum - source_total) < 1e-6,
                "ledger_value_wan_m3": ledger_value if ledger_value is not None else "",
                "ledger_matches_source_total": (
                    abs(ledger_value - source_total) < 1e-6 if ledger_value is not None else False
                ),
            })
    formulas.close()
    values.close()

    by_region = {}
    for region in REGIONS:
        region_rows = {row["date"]: row for row in records if row["region"] == region}
        row24 = region_rows["2024-06-01"]
        row25 = region_rows["2025-05-21"]
        source_ratio = row25["source_total_wan_m3"] / row24["source_total_wan_m3"]
        component_ratio = row25["component_sum_wan_m3"] / row24["component_sum_wan_m3"]
        by_region[region] = {
            "source_total": {
                "2024-06-01": row24["source_total_wan_m3"],
                "2025-05-21": row25["source_total_wan_m3"],
                "ratio_2025_to_2024": source_ratio,
                "direction": "decrease" if source_ratio < 1 else "increase" if source_ratio > 1 else "flat",
            },
            "component_sum": {
                "2024-06-01": row24["component_sum_wan_m3"],
                "2025-05-21": row25["component_sum_wan_m3"],
                "ratio_2025_to_2024": component_ratio,
                "direction": "decrease" if component_ratio < 1 else "increase" if component_ratio > 1 else "flat",
            },
            "direction_robust_to_total_discrepancy": (source_ratio < 1) == (component_ratio < 1),
        }

    output = OUTPUT_ROOT / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output.mkdir(parents=True, exist_ok=False)
    table = output / "manual_region_totals.csv"
    with table.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    report = {
        "status": "complete",
        "experiment": "Cross-region manual pond total and component-sum audit",
        "survey_workbook": str(SURVEY_BOOK),
        "workbook_mtime": datetime.fromtimestamp(SURVEY_BOOK.stat().st_mtime).isoformat(),
        "records": records,
        "regional_trends": by_region,
        "checks": {
            "all_regions_found_on_both_dates": len(records) == 6,
            "all_component_sums_match_source_totals": all(
                row["component_sum_matches_total"] for row in records
            ),
            "all_ledger_values_match_source_totals": all(
                row["ledger_matches_source_total"] for row in records
            ),
            "all_region_directions_robust_to_total_discrepancy": all(
                result["direction_robust_to_total_discrepancy"] for result in by_region.values()
            ),
        },
        "output_table": str(table),
        "interpretation": [
            "Source total rows and recomputed component sums are retained as separate scenarios wherever they differ.",
            "A direction is considered robust only when total-row and component-sum scenarios agree.",
            "The workbook is read only and no source or stage ledger value is changed."
        ],
        "limits": [
            "Survey subunit geometry and exact mapping to model AOIs remain unavailable.",
            "Workbook dates do not establish within-day observation time.",
            "A matching regional direction does not by itself validate model object membership or absolute volume."
        ],
    }
    (output / "summary.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(f"RESULT_FOLDER: {output}")
    print("STATUS: complete")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
